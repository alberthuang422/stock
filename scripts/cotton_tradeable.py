# -*- coding: utf-8 -*-
"""
棉花「可贸易库存」是否比全球总量更能解释棉价？—— 2026-09-12

与白糖（sugar_tradeable.py）同框架。棉花的结构更极端：
  全球库存 69.6 百万 bales 里，中国国储棉 34.7 百万（占 49.8%），2014 年曾占 64.3%。
  中国国储 + 印度 MSP 收储 = 两大"政策缓冲"，不参与或仅边际参与国际贸易。

构造（两条路线）
  [路线A 剥离法] R1 = 全球 − 中国；R2 = 全球 − 中国 − 印度
  [路线B 出口国法] 固定净出口集合（美/巴/澳 + 西非 + 中亚 + 希腊）
      core_sur      = 出口国库存 / 全球消费
      core_sur_own  = 出口国库存 / 出口国自身消费   <- 最贴近"可贸易缓冲"
      core_surplus  = 出口国(产量 − 消费)           [流量，证伪用]
      core_share    = 出口国库存 / 全球库存
      big3_sur      = 美+巴+澳 库存 / 全球消费       <- 三大出口国口径

方法要点（同糖）：
  一阶差分主口径（水平相关被共同趋势污染）；固定集合（动态集合会跳变）。
输出 results/cotton_tradeable.json / data/cotton/cotton_tradeable.csv
"""
import json
import os
import re

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from scipy import stats

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "results", "cotton_tradeable.json")
CSV = os.path.join(BASE, "data", "cotton", "cotton_tradeable.csv")

CORE = ["United States", "Brazil", "Australia", "Greece", "Benin", "Mali", "Burkina Faso",
        "Cote d'Ivoire", "Cameroon", "Chad", "Togo", "Tajikistan", "Kazakhstan",
        "Azerbaijan", "Turkmenistan", "Uzbekistan", "Sudan", "Argentina", "Egypt"]
BIG3 = ["United States", "Brazil", "Australia"]

df = pd.read_csv(os.path.join(BASE, "data", "sugar", "raw", "psd", "psd_alldata.csv"), low_memory=False)
s = df[df.Commodity_Code == 2631000].copy()
s["Value"] = pd.to_numeric(s["Value"], errors="coerce")
piv = s.pivot_table(index=["Country_Name", "Market_Year"], columns="Attribute_Description",
                    values="Value", aggfunc="first").reset_index()
piv["net"] = piv["Exports"].fillna(0) - piv["Imports"].fillna(0)

cov = {}
for y in [1960, 1970, 1980, 1990, 2000, 2010, 2020, 2026]:
    P = piv[piv.Market_Year == y]
    cov[y] = round(float(P[P.Country_Name.isin(CORE)].net.sum() / P[P.net > 0].net.sum() * 100), 1)
print("固定集合净出口覆盖率(%):", cov)

g = pd.read_csv(os.path.join(BASE, "data", "cotton", "cotton_global_sd.csv"))
g = g.rename(columns={g.columns[0]: "MY"}).set_index("MY")
core = piv[piv.Country_Name.isin(CORE)].groupby("Market_Year")[
    ["Ending Stocks", "Production", "Domestic Use", "Exports"]].sum(min_count=1)
big3 = piv[piv.Country_Name.isin(BIG3)].groupby("Market_Year")[
    ["Ending Stocks", "Production", "Domestic Use", "Exports"]].sum(min_count=1)


def cty(name, col):
    return piv[piv.Country_Name == name].set_index("Market_Year")[col].reindex(g.index).fillna(0)


M = pd.DataFrame(index=g.index)
M["w_stocks"] = g["Ending Stocks"]
M["w_sur"] = g["StockToUse"]
M["w_cons"] = g["Domestic Use"]
M["cn"] = cty("China", "Ending Stocks")
M["ind"] = cty("India", "Ending Stocks")
M["R1_stocks"] = M.w_stocks - M.cn
M["R1_sur"] = M.R1_stocks / M.w_cons * 100
M["R2_stocks"] = M.w_stocks - M.cn - M["ind"]
M["R2_sur"] = M.R2_stocks / M.w_cons * 100
M["core_stocks"] = core["Ending Stocks"]
M["core_sur"] = M.core_stocks / M.w_cons * 100
M["core_sur_own"] = M.core_stocks / core["Domestic Use"] * 100
M["core_surplus"] = core["Production"] - core["Domestic Use"]
M["core_share"] = M.core_stocks / M.w_stocks * 100
M["core_exp_share"] = core["Exports"] / g["Exports"] * 100
M["big3_stocks"] = big3["Ending Stocks"]
M["big3_sur"] = M.big3_stocks / M.w_cons * 100

# 价格：MY 期间（Aug Y - Jul Y+1）Pink Sheet 月均（棉花市场年度 Aug-Jul）
wb = load_workbook(os.path.join(BASE, "data", "sugar", "raw", "pink_sheet_2026.xlsx"),
                   read_only=True, data_only=True)
rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
hdr = {str(x).strip(): i for i, x in enumerate(rows[4]) if x}
i_c = hdr["Cotton, A Index"]
PR = {}
for r in rows[6:]:
    if r[0] and re.match(r"^\d{4}M\d{2}$", str(r[0])):
        v = pd.to_numeric(r[i_c], errors="coerce")
        if pd.notna(v):
            PR[str(r[0])] = float(v)


def myavg(y):
    vs = [PR[f"{k//12}M{k%12+1:02d}"] for k in range(y * 12 + 7, (y + 1) * 12 + 7)
          if f"{k//12}M{k%12+1:02d}" in PR]
    return float(np.mean(vs)) if vs else np.nan


M["px"] = [myavg(int(y)) for y in M.index]
M["dlpx"] = np.log(M.px).diff()
M = M.dropna(subset=["px"])

IND = ["w_sur", "w_stocks", "R1_sur", "R1_stocks", "R2_sur", "core_sur", "core_sur_own",
       "core_surplus", "core_share", "big3_sur", "big3_stocks"]
LBL = {"w_sur": "全球库消比", "w_stocks": "全球库存", "R1_sur": "剥离中国·库消比",
       "R1_stocks": "剥离中国·库存", "R2_sur": "剥离中印·库消比",
       "core_sur": "出口国库存/全球消费", "core_sur_own": "出口国库存/自身消费",
       "core_surplus": "可出口盈余（流量）", "core_share": "出口国占全球库存份额",
       "big3_sur": "美巴澳库存/全球消费", "big3_stocks": "美巴澳库存"}


def corr_diff(sub, c):
    d = pd.DataFrame({"x": sub[c].diff(), "y": sub.dlpx}).dropna()
    if len(d) < 10:
        return None
    r, p = stats.pearsonr(d.x, d.y)
    rho, pr = stats.spearmanr(d.x, d.y)
    return {"n": len(d), "r": round(float(r), 3), "R2": round(float(r * r), 3),
            "p": round(float(p), 5), "rho": round(float(rho), 3), "rho_p": round(float(pr), 4)}


full = {c: corr_diff(M, c) for c in IND}
sub1 = {c: corr_diff(M.loc[1960:1999], c) for c in IND}
sub2 = {c: corr_diff(M.loc[2000:2026], c) for c in IND}

print("\n=== [差分] Δ指标 vs Δlog(价格) ===")
print(f"{'口径':26s}{'全样本':>22s}{'1960-99':>14s}{'2000-26':>14s}")
for c in IND:
    f, a, b = full[c], sub1[c], sub2[c]
    print(f"  {LBL[c]:24s} r={f['r']:+.3f} R²={f['R2']:.3f} p={f['p']:.5f} | {a['r']:+.2f} | {b['r']:+.2f}")

# 多元：出口国口径是否含全球口径的增量信息
mv = {}
S = M.loc[2000:2026].copy()
S["dw"] = S.w_sur.diff()
S["dc"] = S.core_sur.diff()
S["dco"] = S.core_sur_own.diff()
S["dr1"] = S.R1_sur.diff()
S["db3"] = S.big3_sur.diff()
d = S[["dw", "dc", "dco", "dr1", "db3", "dlpx"]].dropna()
for xs in [["dw"], ["dc"], ["dco"], ["dr1"], ["db3"], ["dw", "dc"], ["dw", "dco"], ["dw", "db3"]]:
    m = sm.OLS(d.dlpx, sm.add_constant(d[xs])).fit(cov_type="HAC", cov_kwds={"maxlags": 3})
    mv[str(xs)] = {"R2": round(float(m.rsquared), 3),
                   "coef": {k: round(float(m.params[k]), 4) for k in xs},
                   "t": {k: round(float(m.tvalues[k]), 2) for k in xs}}
print("\n=== 2000-2026 多元（HAC）===")
for k, v in mv.items():
    print("  ", k, "R²=", v["R2"], v["coef"], v["t"])

# 当前状态与分位
cur = int(M.index.max())
curstat = {}
for c in IND:
    ser = M[c].dropna()
    v = float(M.loc[cur, c])
    curstat[c] = {"value": round(v, 1), "pct": round(float((ser < v).mean() * 100), 1),
                  "median": round(float(ser.median()), 1),
                  "min": round(float(ser.min()), 1), "min_y": int(ser.idxmin()),
                  "max": round(float(ser.max()), 1), "max_y": int(ser.idxmax())}
print(f"\n=== MY{cur} 当前值与分位 ===")
for c in IND:
    v = curstat[c]
    print(f"  {LBL[c]:24s} {v['value']:>9,.1f}  分位 {v['pct']:5.1f}%  中位 {v['median']:,.1f}")

M.round(2).to_csv(CSV)
out = {
    "meta": {"core_exporters": CORE, "big3": BIG3, "coverage": cov,
             "method": "固定集合；一阶差分主口径；价格=Pink Sheet 棉花 A Index MY 月均(Aug-Jul)",
             "caveat": "水平相关含共同趋势会给出伪正相关，故只用差分"},
    "labels": LBL,
    "diff_full": full, "diff_1960_1999": sub1, "diff_2000_2026": sub2,
    "multivariate_2000_2026": mv,
    "current": curstat,
    "series": {str(y): {c: (None if pd.isna(M.loc[y, c]) else round(float(M.loc[y, c]), 2))
                        for c in IND + ["px", "cn", "ind"]} for y in M.index if y >= 1990},
}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print("\nwrote", OUT, "/", CSV)
