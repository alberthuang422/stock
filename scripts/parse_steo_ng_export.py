# -*- coding: utf-8 -*-
"""解析 EIA STEO（Short-Term Energy Outlook）天然气供需表 → 出口/产量月度序列
- STEO 文件：https://www.eia.gov/outlooks/steo/archives/{mon}{yy}_base.xlsx（如 aug26_base.xlsx）
- 输出 results/steo_ng_export_YYYY.json + .csv（date, lng_export, pipe_export, net_import, dry_prod）
- 5atab 关键行：A20 干气产量 / A23 净进口 / A25 LNG出口 / A27 管道出口
- 日期列：行4 月份(Jan..Dec)，年份按 (col-3)//12 从起始年份偏移（STEO 历史起点通常 2022）
用法: python parse_steo_ng_export.py <xlsx路径> <起始年份> <输出前缀>
"""
import sys
import json
import csv
import openpyxl

MON = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']


def parse(path, start_year=2022, prefix='steo_ng_export'):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['5atab']
    months = {}
    for c in ws[4]:
        if isinstance(c.value, str) and c.value.strip() in MON:
            months[c.column] = c.value.strip()
    dates = []
    for col in sorted(months):
        y = start_year + (col - 3) // 12
        dates.append((col, f'{y}-{months[col]}'))

    def row_vals(r):
        return {d: ws.cell(row=r, column=col).value for col, d in dates}

    lng = row_vals(25)   # LNG gross exports
    pipe = row_vals(27)  # Pipeline gross exports
    neti = row_vals(23)  # Net imports
    dry = row_vals(20)   # Dry gas production

    out = []
    for _, d in dates:
        def f(m):
            v = m.get(d)
            return round(float(v), 2) if isinstance(v, (int, float)) else None
        out.append({'date': d, 'lng_export': f(lng), 'pipe_export': f(pipe),
                    'net_import': f(neti), 'dry_prod': f(dry)})

    root = 'C:/Users/Administrator/Desktop/stock'
    json.dump(out, open(f'{root}/results/{prefix}.json', 'w', encoding='utf-8'),
              ensure_ascii=False, indent=1)
    with open(f'{root}/results/{prefix}.csv', 'w', newline='', encoding='utf-8-sig') as fo:
        w = csv.writer(fo)
        w.writerow(['date', 'lng_export_bcfd', 'pipe_export_bcfd', 'net_import_bcfd', 'dry_prod_bcfd'])
        for r in out:
            w.writerow([r['date'], r['lng_export'], r['pipe_export'], r['net_import'], r['dry_prod']])
    print(f'saved: {len(out)} rows -> results/{prefix}.json/.csv')
    return out


if __name__ == '__main__':
    xlsx = sys.argv[1] if len(sys.argv) > 1 else 'Temp/steo_aug26_base.xlsx'
    sy = int(sys.argv[2]) if len(sys.argv) > 2 else 2022
    pre = sys.argv[3] if len(sys.argv) > 3 else 'steo_ng_export_2026'
    parse(xlsx, sy, pre)