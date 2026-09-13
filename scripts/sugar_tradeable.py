# -*- coding: utf-8 -*-
"""
白糖「可贸易库存」是否比全球总量更能解释糖价？—— 2026-09-12

动机
  全球期末库存 44.4 Mt 里，中国 402 万吨基本不参与国际贸易、印度 651 万吨受出口管制，
  巴西只剩 22 万吨。直觉上"可贸易口径"应比总量更贴近价格。本脚本做实测。

构造（两条路线，互为稳健性）
  [路线A 剥离法] 固定国家集，从全球库存中剥离封闭库存
      R1 = 全球 − 中国          R2 = 全球 − 中国 − 印度
  [路线B 出口国法] 固定 10 国出口集合（净出口口径，排除 UAE/沙特等再出口国）
      core_sur      = 出口国库存 / 全球消费
      core_sur_own  = 出口国库存 / 出口国自身消费     <- 最贴近"可贸易缓冲"
      core_surplus  = 出口国(产量 − 消费)             [流量，用于证伪]
      core_share    = 出口国库存 / 全球库存

方法要点（踩过的坑）
  ⚠️ 水平相关会被**共同趋势**污染：库存与价格都长期上行 → 出现 r=+0.42 的伪正相关。
     必须用**一阶差分**做主口径（Δ指标 vs Δlog价格）。
  ⚠️ 动态集合（每年取覆盖 80% 的国家）会因成员切换产生跳变（2018 年 T1 跳 3.3 倍），
     把差分信号淹没 → 必须用**固定集合**。
输出 results/sugar_tradeable.json
"""
import json
import os
import re

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from scipy import stats

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "results", "sugar_tradeable.json")
CSV = os.path.join(BASE, "data", "sugar", "sugar_tradeable.csv")

CORE = ["Brazil", "Thailand", "Australia", "India", "Guatemala", "Mexico",
        "South Africa", "Colombia", "Ukraine", "Argentina"]

df = pd.read_csv(os.path.join(BASE, "data", "sugar", "raw", "psd", "psd_alldata.csv"), low_memory=False)
s = df[df.Commodity_Code == 612000].copy()
s["Value"] = pd.to_numeric(s["Value"], errors="coerce")
piv = s.pivot_table(index=["Country_Name", "Market_Year"], columns="Attribute_Description",
                    values="Value", aggfunc="first").reset_index()
piv["net"] = piv["Exports"].fillna(0) - piv["Imports"].fillna(0)

cov = {}
for y in [1960, 1970, 1980, 1990, 2000, 2010, 2020, 2026]:
    P = piv[piv.Market_Year == y]
    cov[y] = round(float(P[P.Country_Name.isin(CORE)].net.sum() / P[P.net > 0].net.sum() * 100), 1)
print("固定集合净出口覆盖率:", cov)

g = pd.read_csv(os.path.join(BASE, "data", "sugar", "sugar_global_sd.csv"))
g = g.rename(columns={g.columns[0]: "MY"}).set_index("MY")
core = piv[piv.Country_Name.isin(CORE)].groupby("Market_Year")[
    ["Ending Stocks", "Production", "Total Disappearance", "Exports"]].sum(min_count=1)


def cty(name, col):
    return piv[piv.Country_Name == name].set_index("Market_Year")[col].reindex(g.index).fillna(0)


M = pd.DataFrame(index=g.index)
M["w_stocks"] = g["Ending Stocks"]
M["w_sur"] = g["StockToUse"]
M["w_cons"] = g["Total Disappearance"]
M["cn"] = cty("China", "Ending Stocks")
M["ind"] = cty("India", "Ending Stocks")
M["R1_stocks"] = M.w_stocks - M.cn
M["R1_sur"] = M.R1_stocks / M.w_cons * 100
M["R2_stocks"] = M.w_stocks - M.cn - M["ind"]
M["R2_sur"] = M.R2_stocks / M.w_cons * 100
M["core_stocks"] = core["Ending Stocks"]
M["core_sur"] = M.core_stocks / M.w_cons * 100
M["core_sur_own"] = M.core_stocks / core["Total Disappearance"] * 100
M["core_surplus"] = core["Production"] - core["Total Disappearance"]
M["core_share"] = M.core_stocks / M.w_stocks * 100
M["core_exp_share"] = core["Exports"] / g["Exports"] * 100

# 价格：MY 期间（Oct Y - Sep Y+1）Pink Sheet 月均
wb = load_workbook(os.path.join(BASE, "data", "sugar", "raw", "pink_sheet_2026.xlsx"),
                   read_only=True, data_only=True)
rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
hdr = {str(x).strip(): i for i, x in enumerate(rows[4]) if x}
i_s = hdr["Sugar, world"]
PR = {}
for r in rows[6:]:
    if r[0] and re.match(r"^\d{4}M\d{2}$", str(r[0])):
        v = pd.to_numeric(r[i_s], errors="coerce")
        if pd.notna(v):
            PR[str(r[0])] = float(v)


def myavg(y):
    vs = [PR[f"{k//12}M{k%12+1:02d}"] for k in range(y * 12 + 9, (y + 1) * 12 + 9)
          if f"{k//12}M{k%12+1:02d}" in PR]
    return float(np.mean(vs)) if vs else np.nan


M["px"] = [myavg(int(y)) for y in M.index]
M["dlpx"] = np.log(M.px).diff()
M = M.dropna(subset=["px"])

IND = ["w_sur", "w_stocks", "R1_sur", "R1_stocks", "R2_sur", "core_sur", "core_sur_own",
       "core_surplus", "core_share"]
LBL = {"w_sur": "全球库消比", "w_stocks": "全球库存", "R1_sur": "剥离中国·库消比",
       "R1_stocks": "剥离中国·库存", "R2_sur": "剥离中印·库消比",
       "core_sur": "出口国库存/全球消费", "core_sur_own": "出口国库存/自身消费",
       "core_surplus": "可出口盈余（流量）", "core_share": "出口国占全球库存份额"}


def corr_diff(sub, c):
    d = pd.DataFrame({"x": sub[c].diff(), "y": sub.dlpx}).dropna()
    if len(d) < 10:
        return None
    r, p = stats.pearsonr(d.x, d.y)
    rho, pr = stats.spearmanr(d.x, d.y)
    return {"n": len(d), "r": round(float(r), 3), "R2": round(float(r * r), 3),
            "p": round(float(p), 5), "rho": round(float(rho), 3), "rho_p": round(float(pr), 4)}


full = {c: corr_diff(M, c) for c in IND}
sub1 = {c: corr_diff(M.loc[1960:1999], c) for c in IND}
sub2 = {c: corr_diff(M.loc[2000:2026], c) for c in IND}

print("\n=== [差分] Δ指标 vs Δlog(价格) ===")
print(f"{'口径':26s}{'全样本':>22s}{'1960-99':>14s}{'2000-26':>14s}")
for c in IND:
    f, a, b = full[c], sub1[c], sub2[c]
    print(f"  {LBL[c]:24s} r={f['r']:+.3f} R²={f['R2']:.3f} p={f['p']:.5f} | {a['r']:+.2f} | {b['r']:+.2f}")

# 多元：出口国口径是否含全球口径的增量信息
mv = {}
S = M.loc[2000:2026].copy()
S["dw"] = S.w_sur.diff()
S["dc"] = S.core_sur.diff()
S["dco"] = S.core_sur_own.diff()
d = S[["dw", "dc", "dco", "dlpx"]].dropna()
for xs in [["dw"], ["dc"], ["dco"], ["dw", "dc"], ["dw", "dco"]]:
    m = sm.OLS(d.dlpx, sm.add_constant(d[xs])).fit(cov_type="HAC", cov_kwds={"maxlags": 3})
    mv[str(xs)] = {"R2": round(float(m.rsquared), 3),
                   "coef": {k: round(float(m.params[k]), 4) for k in xs},
                   "t": {k: round(float(m.tvalues[k]), 2) for k in xs}}
print("\n=== 2000-2026 多元（HAC）===")
for k, v in mv.items():
    print("  ", k, "R²=", v["R2"], v["coef"], v["t"])

# 当前状态与分位
cur = int(M.index.max())
curstat = {}
for c in IND:
    ser = M[c].dropna()
    v = float(M.loc[cur, c])
    curstat[c] = {"value": round(v, 1), "pct": round(float((ser < v).mean() * 100), 1),
                  "median": round(float(ser.median()), 1),
                  "min": round(float(ser.min()), 1), "min_y": int(ser.idxmin()),
                  "max": round(float(ser.max()), 1), "max_y": int(ser.idxmax())}
print(f"\n=== MY{cur} 当前值与分位 ===")
for c in IND:
    v = curstat[c]
    print(f"  {LBL[c]:24s} {v['value']:>9,.1f}  分位 {v['pct']:5.1f}%  中位 {v['median']:,.1f}")

last = int(M.index.max() - 1)
M.round(2).to_csv(CSV)
out = {
    "meta": {"core_exporters": CORE, "coverage": cov,
             "method": "固定集合；一阶差分主口径；价格=Pink Sheet 世界糖价 MY 月均(Oct-Sep)",
             "caveat": "水平相关含共同趋势会给出伪正相关，故只用差分"},
    "labels": LBL,
    "diff_full": full, "diff_1960_1999": sub1, "diff_2000_2026": sub2,
    "multivariate_2000_2026": mv,
    "current": curstat,
    "series": {str(y): {c: (None if pd.isna(M.loc[y, c]) else round(float(M.loc[y, c]), 2))
                        for c in IND + ["px", "cn", "ind"]} for y in M.index if y >= 1990},
}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print("\nwrote", OUT, "/", CSV)
